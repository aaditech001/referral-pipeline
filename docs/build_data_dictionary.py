"""Builds docs/data_dictionary.xlsx (business-friendly) - re-run if the report changes."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

COLS = [
 ("referral_details_id","Reference number of the most recent activity record for this referral.","Whole number","0 means no activity record exists yet for this referral.","101"),
 ("referral_id","Unique ID of one referral (one person recommending another). Each row of the report is one referral.","Text","Never empty, never repeated.","9331c8f144dad5a3b8e4a10467b4343a"),
 ("referral_source","How the referral was made.","Text","User Sign Up, Draft Transaction or Lead.","User Sign Up"),
 ("referral_source_category","Whether the referral came from an online or offline channel.","Text","Online or Offline. For Lead referrals it comes from the lead's own channel; Unknown if the lead cannot be found.","Online"),
 ("referral_at","Date and time the referral was created, in the referrer's local time.","Date & time","Local time of the referrer's home club.","2024-05-01 12:17:31"),
 ("referrer_id","ID of the existing member who made the referral.","Text","Unknown if the referral has no referrer.","2c71c5d66c7e12a0b3c200ba6ed3b78e"),
 ("referrer_name","Name of the existing member who referred.","Text","Unknown if the member is not found.","John Doe"),
 ("referrer_phone_number","Phone number of the referrer.","Text","Unknown if the member is not found.","123-456-7890"),
 ("referrer_homeclub","Club the referrer belongs to.","Text","Club names are shown in capitals.","PERMATA HIJAU"),
 ("referee_id","ID of the referred person. Only filled for Lead referrals.","Text","Unknown for non-Lead referrals.","f12348hbsdkjkfhkjdf"),
 ("referee_name","Name of the person who was referred.","Text","Unknown if not provided.","Jhon Doe"),
 ("referee_phone","Phone number of the person who was referred.","Text","","987-654-3210"),
 ("referral_status","Current outcome of the referral.","Text","Berhasil (successful), Menunggu (pending) or Tidak Berhasil (failed).","Berhasil"),
 ("num_reward_days","Number of free membership days promised as the reward.","Whole number","0 means no reward was assigned.","30"),
 ("transaction_id","ID of the purchase linked to the referral.","Text","Unknown if there is no linked purchase.","1d1eb8a9e864a1cccb2d850398461807"),
 ("transaction_status","Payment status of the linked purchase.","Text","Paid, or Unknown if no purchase details exist.","Paid"),
 ("transaction_at","Date and time of the purchase, in the local time of the place where it happened.","Date & time","Blank if there is no purchase.","2024-05-20 11:23:09"),
 ("transaction_location","Club or place where the purchase happened.","Text","Shown in capitals; Unknown if no purchase.","BENHIL"),
 ("transaction_type","Kind of purchase.","Text","New (first-time membership), Rejoin (returning member), or Unknown.","New"),
 ("updated_at","Date and time the referral record was last changed, in the referrer's local time.","Date & time","","2024-05-02 09:00:00"),
 ("reward_granted_at","Date and time the reward was given.","Date & time","Blank if the reward has not been granted.","2024-05-30 14:00:00"),
 ("is_business_logic_valid","Whether the referral and its reward pass all business checks.","True / False","TRUE = valid. FALSE = invalid / needs review (see 'Validity Rules' sheet).","TRUE"),
]
RULES = [
 ("VALID","A reward is valid when ALL are true","Reward is above 0 days; status is Berhasil; there is a linked purchase; it is Paid and of type New; it happened after the referral and in the same month; the referrer's membership has not expired and the account is not deleted; the reward was granted."),
 ("VALID","No reward expected","Status is Menunggu (pending) or Tidak Berhasil (failed) AND no reward is assigned."),
 ("INVALID","Reward on an unsuccessful referral","Reward above 0 days but status is not Berhasil."),
 ("INVALID","Reward without a purchase","Reward above 0 days but no linked purchase."),
 ("INVALID","Paid purchase but no reward assigned","No reward, but there is a Paid purchase made after the referral."),
 ("INVALID","Successful but no reward","Status is Berhasil but the reward is empty or 0."),
 ("INVALID","Purchase before referral","The purchase happened before the referral was created."),
]
NOTES = [
 "A referral is flagged FALSE if it matches ANY invalid rule, or if it does not match any valid rule.",
 "Times are local: converted from UTC using the timezone of the referrer's home club (referral, update and reward times) or of the purchase location (purchase time).",
 "Where a value does not exist the report shows 'Unknown' (text) or 0 (numbers). Date columns are left blank instead of inventing a date.",
 "One row = one referral. The report contains 46 rows.",
]

F = "Arial"
head_fill = PatternFill("solid", start_color="1F4E78")
thin = Side(style="thin", color="BBBBBB"); box = Border(left=thin,right=thin,top=thin,bottom=thin)

def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for r in rows: ws.append(list(r))
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10); c.border = box
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(widths): ws.column_dimensions[chr(65+i)].width = w
    ws.freeze_panes = "A2"

wb = Workbook()
ws = wb.active; ws.title = "Report Columns"
sheet(ws, ["Column","What it means","Type","Allowed values / notes","Example"], COLS, [26,55,14,50,34])
ws2 = wb.create_sheet("Validity Rules")
sheet(ws2, ["Result","Rule","Plain-English explanation"], RULES, [12,38,100])
ws3 = wb.create_sheet("Notes")
sheet(ws3, ["General notes"], [(n,) for n in NOTES], [140])
wb.save("docs/data_dictionary.xlsx")
